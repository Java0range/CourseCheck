from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Length

from schemas import CourseworkCheckResult
from settings import (
    FORMATTING_REQUIREMENTS,
    LMSTUDIO_API_KEY,
    LMSTUDIO_BASE_URL,
    LMSTUDIO_MODEL,
    MAIN_CONTENT_START_PATTERNS,
    MAX_COMPLETION_TOKENS,
    MAX_PROMPT_CHARS,
    REQUEST_TIMEOUT_SECONDS,
    SKIP_BEFORE_MAIN_CONTENT,
    SUBMITTED_MARKER,
    TEMPERATURE,
)


@dataclass
class ParagraphInfo:
    paragraph_id: int
    text: str
    is_empty: bool
    role_hint: str
    heading_level_hint: Optional[int]
    style_name: Optional[str]
    alignment: Optional[str]
    first_line_indent_cm: Optional[float]
    line_spacing: Optional[Any]
    space_before_pt: Optional[float]
    space_after_pt: Optional[float]
    font_name_most_common: Optional[str]
    font_size_pt_most_common: Optional[float]
    prev_is_empty: Optional[bool]
    next_is_empty: Optional[bool]


@dataclass
class DocumentExtract:
    file_name: str
    paragraphs: list[ParagraphInfo]
    skipped_before_paragraph_id: Optional[int]

    @property
    def checked_paragraphs(self) -> list[ParagraphInfo]:
        if self.skipped_before_paragraph_id is None:
            return self.paragraphs
        return [
            paragraph
            for paragraph in self.paragraphs
            if paragraph.paragraph_id >= self.skipped_before_paragraph_id
        ]


# =========================
# DOCX helpers
# =========================


def length_to_cm(value: Optional[Length]) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(value.cm, 2)
    except Exception:
        return None


def length_to_pt(value: Optional[Length]) -> Optional[float]:
    if value is None:
        return None
    try:
        return round(value.pt, 2)
    except Exception:
        return None


def alignment_to_str(value: Any) -> Optional[str]:
    if value is None:
        return None

    mapping = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribute",
    }
    return mapping.get(value, str(value))


def line_spacing_to_value(value: Any) -> Optional[Any]:
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return round(float(value), 2)

    try:
        if isinstance(value, Length):
            return {"pt": round(value.pt, 2)}
    except Exception:
        pass

    return str(value)


def safe_style_name(paragraph: Any) -> Optional[str]:
    try:
        if paragraph.style is not None:
            return paragraph.style.name
    except Exception:
        return None
    return None


def get_paragraph_style_font(paragraph: Any) -> tuple[Optional[str], Optional[float]]:
    try:
        style = paragraph.style
        if style is None:
            return None, None

        font_name = style.font.name
        font_size = style.font.size.pt if style.font.size is not None else None
        return font_name, round(font_size, 2) if font_size else None
    except Exception:
        return None, None


def get_run_font_name(run: Any, paragraph_style_font_name: Optional[str]) -> Optional[str]:
    try:
        if run.font is not None and run.font.name:
            return run.font.name
    except Exception:
        pass
    return paragraph_style_font_name


def get_run_font_size(run: Any, paragraph_style_font_size: Optional[float]) -> Optional[float]:
    try:
        if run.font is not None and run.font.size is not None:
            return round(run.font.size.pt, 2)
    except Exception:
        pass
    return paragraph_style_font_size


def most_common_or_none(values: list[Any]) -> Optional[Any]:
    filtered = [value for value in values if value is not None]
    if not filtered:
        return None
    return Counter(filtered).most_common(1)[0][0]


def detect_heading(text: str, style_name: Optional[str]) -> tuple[str, Optional[int]]:
    stripped = text.strip()
    lower_style = (style_name or "").lower()

    if not stripped:
        return "empty", None

    style_heading_match = re.search(r"(heading|заголовок)\s*(\d+)?", lower_style)
    if style_heading_match:
        level_text = style_heading_match.group(2)
        level = int(level_text) if level_text else 1
        if level == 1:
            return "chapter_heading", 1
        return "subheading", level

    upper = stripped.upper()
    major_headings = {
        "ВВЕДЕНИЕ",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК ЛИТЕРАТУРЫ",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        "СПИСОК ИСПОЛЬЗУЕМОЙ ЛИТЕРАТУРЫ",
        "ПРИЛОЖЕНИЯ",
        "ПРИЛОЖЕНИЕ",
        "СОДЕРЖАНИЕ",
        "ОГЛАВЛЕНИЕ",
    }

    if upper in major_headings:
        return "chapter_heading", 1

    if re.match(r"^(ГЛАВА|РАЗДЕЛ)\s+\d+", upper):
        return "chapter_heading", 1

    if re.match(r"^\d+\.\s+\S+", stripped):
        return "chapter_heading", 1

    if re.match(r"^\d+\.\d+(\.\d+)*\.?\s+\S+", stripped):
        level = stripped.split()[0].count(".") + 1
        return "subheading", max(level, 2)

    if len(stripped) <= 100 and not stripped.endswith((".", ",", ";", ":")):
        if sum(ch.isalpha() for ch in stripped) >= 5:
            return "possible_heading", None

    return "body", None


def looks_like_main_content_start(text: str) -> bool:
    stripped = text.strip().upper()
    if not stripped:
        return False

    for pattern in MAIN_CONTENT_START_PATTERNS:
        if re.match(pattern, stripped, flags=re.IGNORECASE):
            return True
    return False


def find_main_content_start(paragraphs: list[ParagraphInfo]) -> Optional[int]:
    if not SKIP_BEFORE_MAIN_CONTENT:
        return None

    for paragraph in paragraphs:
        if looks_like_main_content_start(paragraph.text):
            return paragraph.paragraph_id

    return None


def extract_docx(file_path: Path) -> DocumentExtract:
    if not file_path.exists():
        raise FileNotFoundError(f"Файл не найден: {file_path}")

    if file_path.suffix.lower() != ".docx":
        raise ValueError("ИИ-проверка сейчас поддерживает только .docx")

    document = Document(str(file_path))
    paragraphs: list[ParagraphInfo] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text or ""
        stripped = text.strip()
        is_empty = len(stripped) == 0

        style_name = safe_style_name(paragraph)
        role_hint, heading_level_hint = detect_heading(stripped, style_name)
        paragraph_format = paragraph.paragraph_format
        paragraph_style_font_name, paragraph_style_font_size = get_paragraph_style_font(paragraph)

        run_font_names: list[Optional[str]] = []
        run_font_sizes: list[Optional[float]] = []

        for run in paragraph.runs:
            if not run.text:
                continue
            run_font_names.append(get_run_font_name(run, paragraph_style_font_name))
            run_font_sizes.append(get_run_font_size(run, paragraph_style_font_size))

        paragraphs.append(
            ParagraphInfo(
                paragraph_id=index,
                text=text,
                is_empty=is_empty,
                role_hint=role_hint,
                heading_level_hint=heading_level_hint,
                style_name=style_name,
                alignment=alignment_to_str(paragraph.alignment),
                first_line_indent_cm=length_to_cm(paragraph_format.first_line_indent),
                line_spacing=line_spacing_to_value(paragraph_format.line_spacing),
                space_before_pt=length_to_pt(paragraph_format.space_before),
                space_after_pt=length_to_pt(paragraph_format.space_after),
                font_name_most_common=most_common_or_none(run_font_names),
                font_size_pt_most_common=most_common_or_none(run_font_sizes),
                prev_is_empty=None,
                next_is_empty=None,
            )
        )

    for index, paragraph in enumerate(paragraphs):
        paragraph.prev_is_empty = paragraphs[index - 1].is_empty if index > 0 else None
        paragraph.next_is_empty = paragraphs[index + 1].is_empty if index < len(paragraphs) - 1 else None

    return DocumentExtract(
        file_name=file_path.name,
        paragraphs=paragraphs,
        skipped_before_paragraph_id=find_main_content_start(paragraphs),
    )


# =========================
# Local formatting check
# =========================


def almost_equal(value: Optional[float], expected: float, tolerance: float = 0.08) -> bool:
    if value is None:
        return True
    return abs(value - expected) <= tolerance


def font_is_times_new_Roman(font_name: Optional[str]) -> bool:
    if font_name is None:
        return True
    normalized = font_name.lower().replace(" ", "")
    return "timesnewroman" in normalized


def line_spacing_is_15(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, (int, float)):
        return abs(float(value) - 1.5) <= 0.08
    if isinstance(value, dict) and "pt" in value:
        return True
    return True


def check_local_formatting(document: DocumentExtract) -> list[str]:
    issues: list[str] = []

    checked_paragraphs = document.checked_paragraphs
    if document.skipped_before_paragraph_id is None and SKIP_BEFORE_MAIN_CONTENT:
        issues.append(
            "Не найден явный маркер начала основной части (например, 'Введение' или 'Глава 1'), "
            "поэтому титульный лист мог попасть в проверку."
        )

    for paragraph in checked_paragraphs:
        if paragraph.is_empty:
            continue

        paragraph_ref = f"Абзац {paragraph.paragraph_id}"

        if paragraph.role_hint == "body":
            if not font_is_times_new_Roman(paragraph.font_name_most_common):
                issues.append(
                    f"{paragraph_ref}: основной текст должен быть Times New Roman, "
                    f"обнаружено: {paragraph.font_name_most_common}."
                )

            if paragraph.font_size_pt_most_common is not None and not almost_equal(
                paragraph.font_size_pt_most_common,
                14,
            ):
                issues.append(
                    f"{paragraph_ref}: размер основного текста должен быть 14 pt, "
                    f"обнаружено: {paragraph.font_size_pt_most_common} pt."
                )

            if paragraph.alignment is not None and paragraph.alignment != "justify":
                issues.append(
                    f"{paragraph_ref}: основной текст должен быть выровнен по ширине, "
                    f"обнаружено: {paragraph.alignment}."
                )

            if paragraph.first_line_indent_cm is not None and not almost_equal(
                paragraph.first_line_indent_cm,
                1.25,
                tolerance=0.12,
            ):
                issues.append(
                    f"{paragraph_ref}: красная строка должна быть 1.25 см, "
                    f"обнаружено: {paragraph.first_line_indent_cm} см."
                )

            if not line_spacing_is_15(paragraph.line_spacing):
                issues.append(
                    f"{paragraph_ref}: межстрочный интервал должен быть 1.5, "
                    f"обнаружено: {paragraph.line_spacing}."
                )

        elif paragraph.role_hint == "chapter_heading":
            if not font_is_times_new_Roman(paragraph.font_name_most_common):
                issues.append(
                    f"{paragraph_ref}: заголовок главы должен быть Times New Roman, "
                    f"обнаружено: {paragraph.font_name_most_common}."
                )

            if paragraph.font_size_pt_most_common is not None and not almost_equal(
                paragraph.font_size_pt_most_common,
                14,
            ):
                issues.append(
                    f"{paragraph_ref}: размер заголовка главы должен быть 14 pt, "
                    f"обнаружено: {paragraph.font_size_pt_most_common} pt."
                )

            if paragraph.alignment is not None and paragraph.alignment != "center":
                issues.append(
                    f"{paragraph_ref}: заголовок главы должен быть выровнен по центру, "
                    f"обнаружено: {paragraph.alignment}."
                )

            if paragraph.next_is_empty is False:
                issues.append(
                    f"{paragraph_ref}: после заголовка главы должна быть пустая строка."
                )

        elif paragraph.role_hint in {"subheading", "possible_heading"}:
            if not font_is_times_new_Roman(paragraph.font_name_most_common):
                issues.append(
                    f"{paragraph_ref}: заголовок должен быть Times New Roman, "
                    f"обнаружено: {paragraph.font_name_most_common}."
                )

            if paragraph.font_size_pt_most_common is not None and not almost_equal(
                paragraph.font_size_pt_most_common,
                14,
            ):
                issues.append(
                    f"{paragraph_ref}: размер заголовка должен быть 14 pt, "
                    f"обнаружено: {paragraph.font_size_pt_most_common} pt."
                )

            if paragraph.alignment is not None and paragraph.alignment != "left":
                issues.append(
                    f"{paragraph_ref}: заголовок должен быть выровнен по левому краю, "
                    f"обнаружено: {paragraph.alignment}."
                )

            if paragraph.prev_is_empty is False:
                issues.append(f"{paragraph_ref}: перед заголовком должна быть пустая строка.")
            if paragraph.next_is_empty is False:
                issues.append(f"{paragraph_ref}: после заголовка должна быть пустая строка.")

    return issues


# =========================
# Prompt and LM Studio
# =========================


def paragraph_to_text(paragraph: ParagraphInfo) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    return f"[{paragraph.paragraph_id}] {text}"


def build_prompt(
    *,
    topic: str,
    coursework_desc: str,
    document: DocumentExtract,
    local_issues: list[str],
) -> tuple[str, bool]:
    checked_text = "\n".join(
        item
        for item in (paragraph_to_text(paragraph) for paragraph in document.checked_paragraphs)
        if item
    )

    prompt_was_truncated = False
    if len(checked_text) > MAX_PROMPT_CHARS:
        checked_text = checked_text[:MAX_PROMPT_CHARS]
        prompt_was_truncated = True

    skipped_note = ""
    if document.skipped_before_paragraph_id is not None:
        skipped_note = (
            "Титульный лист и служебные абзацы до начала основной части не проверяются. "
            f"Проверка начинается с абзаца {document.skipped_before_paragraph_id}."
        )
    elif SKIP_BEFORE_MAIN_CONTENT:
        skipped_note = (
            "Явный маркер начала основной части не найден, поэтому документ проверяется с первого абзаца."
        )

    local_issues_text = "\n".join(f"- {issue}" for issue in local_issues)
    if not local_issues_text:
        local_issues_text = "Явных локальных нарушений оформления не найдено."

    desc_text = coursework_desc.strip() or "Описание курсовой не задано."

    prompt = f"""
Ты — строгий эксперт по проверке курсовых работ.

Нужно проверить работу на соответствие теме и требованиям оформления.

Тема курсовой:
{topic}

Описание курсовой:
{desc_text}

Требования к оформлению:
{FORMATTING_REQUIREMENTS}

Правила проверки:
- Не возвращай JSON.
- Напиши обычный комментарий для студента на русском языке.
- Если есть замечания, явно перечисли, что нужно исправить.
- Если работа не соответствует теме, напиши это прямо.
- Если данных недостаточно, напиши, каких данных недостаточно.
- Не выдумывай нарушения.
- Ссылайся на номера абзацев в квадратных скобках, если замечание относится к конкретному месту.
- {skipped_note}
- Если есть хотя бы одно замечание, НЕ пиши маркер {SUBMITTED_MARKER}.
- Если замечаний нет вообще, в самом конце ответа добавь отдельной строкой только {SUBMITTED_MARKER}.

Локальная проверка оформления уже нашла:
{local_issues_text}

Текст проверяемой части документа:
{checked_text}
""".strip()

    return prompt, prompt_was_truncated


async def list_lmstudio_models() -> str:
    url = f"{LMSTUDIO_BASE_URL.rstrip('/')}/models"
    headers = {"Authorization": f"Bearer {LMSTUDIO_API_KEY}"}

    try:
        async with httpx.AsyncClient(timeout=REQUEST_TIMEOUT_SECONDS) as client:
            response = await client.get(url, headers=headers)
            response.raise_for_status()
            data = response.json()
    except Exception as exc:
        return f"Не удалось получить список моделей LM Studio: {exc}"

    models = data.get("data", [])
    if not models:
        return "LM Studio не вернул список моделей."

    return "Доступные модели LM Studio:\n" + "\n".join(
        f"- {model.get('id', '<unknown>')}" for model in models
    )


async def ask_lmstudio(prompt: str) -> str:
    url = f"{LMSTUDIO_BASE_URL.rstrip('/')}/chat/completions"

    payload = {
        "model": LMSTUDIO_MODEL,
        "messages": [
            {
                "role": "system",
                "content": (
                    "Ты проверяешь курсовые работы. Отвечай только обычным текстом, "
                    "без JSON и без Markdown-таблиц."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        "temperature": TEMPERATURE,
        "max_tokens": MAX_COMPLETION_TOKENS,
    }

    headers = {
        "Authorization": f"Bearer {LMSTUDIO_API_KEY}",
        "Content-Type": "application/json",
    }

    async with httpx.AsyncClient(timeout=REQUEST_TIMEOUT_SECONDS) as client:
        response = await client.post(url, json=payload, headers=headers)

    if response.status_code >= 400:
        server_text = response.text.strip()
        models_text = await list_lmstudio_models()
        raise RuntimeError(
            f"LM Studio вернул ошибку {response.status_code}.\n\n"
            f"Ответ сервера:\n{server_text}\n\n{models_text}"
        )

    data = response.json()
    try:
        return str(data["choices"][0]["message"]["content"]).strip()
    except Exception as exc:
        raise RuntimeError(f"Не удалось прочитать ответ LM Studio: {data}") from exc


# =========================
# Result normalization
# =========================


def remove_submitted_marker(comment: str) -> str:
    return comment.replace(SUBMITTED_MARKER, "").strip()


def llm_comment_claims_success(comment: str) -> bool:
    lines = [line.strip() for line in comment.splitlines() if line.strip()]
    return bool(lines) and lines[-1] == SUBMITTED_MARKER


def llm_comment_contains_issue_words(comment: str) -> bool:
    clean = remove_submitted_marker(comment).lower()

    positive_phrases = (
        "замечаний нет",
        "нарушений не обнаружено",
        "исправления не требуются",
        "работа соответствует",
    )
    if any(phrase in clean for phrase in positive_phrases):
        negative_without_context = clean
        for phrase in positive_phrases:
            negative_without_context = negative_without_context.replace(phrase, "")
    else:
        negative_without_context = clean

    issue_markers = (
        "нужно исправить",
        "следует исправить",
        "замечание",
        "замечания",
        "нарушение",
        "нарушения",
        "не соответствует",
        "ошибка",
        "ошибки",
        "отсутствует",
        "недостаточно данных",
        "не удалось проверить",
        "неверно",
        "проблема",
    )

    return any(marker in negative_without_context for marker in issue_markers)


def normalize_comment(
    *,
    raw_comment: str,
    local_issues: list[str],
    prompt_was_truncated: bool,
) -> CourseworkCheckResult:
    comment = raw_comment.strip()

    if not comment:
        comment = "ИИ не вернул комментарий. Требуется ручная проверка преподавателем."

    enforced_issues = list(local_issues)
    if prompt_was_truncated:
        enforced_issues.append(
            "Документ слишком большой для одного запроса LM Studio, поэтому проверка выполнена не полностью."
        )

    if enforced_issues:
        cleaned_ai_comment = remove_submitted_marker(comment)
        local_text = "\n".join(f"- {issue}" for issue in enforced_issues)
        final_comment = (
            "Автоматическая проверка нашла замечания:\n"
            f"{local_text}\n\n"
            "Комментарий ИИ:\n"
            f"{cleaned_ai_comment}"
        ).strip()
        return CourseworkCheckResult(accepted=False, comment=final_comment)

    if llm_comment_claims_success(comment) and not llm_comment_contains_issue_words(comment):
        final_comment = remove_submitted_marker(comment)
        if final_comment:
            final_comment = f"{final_comment}\n{SUBMITTED_MARKER}"
        else:
            final_comment = SUBMITTED_MARKER
        return CourseworkCheckResult(accepted=True, comment=final_comment)

    cleaned_comment = remove_submitted_marker(comment)
    return CourseworkCheckResult(accepted=False, comment=cleaned_comment)


async def check_coursework(
    *,
    file_path: Path,
    topic: str,
    coursework_desc: str = "",
) -> CourseworkCheckResult:
    document = extract_docx(file_path)
    local_issues = check_local_formatting(document)

    prompt, prompt_was_truncated = build_prompt(
        topic=topic,
        coursework_desc=coursework_desc,
        document=document,
        local_issues=local_issues,
    )

    raw_comment = await ask_lmstudio(prompt)

    return normalize_comment(
        raw_comment=raw_comment,
        local_issues=local_issues,
        prompt_was_truncated=prompt_was_truncated,
    )
